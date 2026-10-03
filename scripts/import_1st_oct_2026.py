"""Import the "1st oct 2026" tab of the Murshid Store sheet into the Bahria branch.

    Server:  bash scripts/run_import_1st_oct_2026.sh           (check -- writes nothing)
             APPLY=1 bash scripts/run_import_1st_oct_2026.sh   (create / update)
    Local:   IMPORT_BRANCH="Branch 2" venv/Scripts/python.exe odoo/odoo-bin shell \
                 -c odoo.conf -d OStore --no-http < custom_addons/pos_retail/scripts/import_1st_oct_2026.py

ONLY this tab is read, and ONLY the products this script created are ever
updated. The "Bahria ..." tabs and their products are not read, compared or
touched -- that is import_catalogue.py's job -- and import_catalogue.py in turn
leaves these products alone: they carry their own ids (tab_1st_oct_2026_...),
which it does not recognise as its own, so it never archives them as "no
longer in the sheet".

WITHOUT APPLY=1 THIS IS A CHECK: it prints the products it would create, every
field it would change on the ones it created before, the opening stock it
would load, and the rows it skips. Nothing is written.

How a row becomes a product -- the same rules as the Bahria import:
  * the department comes from the section heading above the row (a name on
    its own after an empty row): Sanitary/plumbering, PAINT& CHEMICALS,
    ELECTRIC, HARDWARE; the Product Category column wins when it is filled.
  * name exactly as in the sheet; size into the Size field; brand from the
    first Brand column; sold at the Bahria branch only; cost in every company.
  * identity: name + size + brand within this tab. Rows that repeat one
    product are merged (quantities added); a row repeating an earlier row's
    name, price and quantity is skipped as a copy.
  * a quantity or price with a unit ("50 ft", "150 kg") makes the product in
    that unit; "pcs" and "set" are pieces.
  * opening stock = the quantity, loaded ONCE per product; after that sales
    and adjustments keep it right and the sheet never overwrites it.
  * no sales price: imported at 0 and listed -- the till refuses to sell it
    until the sheet has a price.

Environment (all optional):
    APPLY=1          write the changes
    IMPORT_BRANCH    branch that sells and stocks the products (default Bahria)
    IMPORT_TAB       another tab of the same layout (default "1st oct 2026")
    IMPORT_XLSX      a local .xlsx instead of the Google Sheet
"""

import datetime
import io
import os
import re
import urllib.request

import openpyxl

APPLY = os.environ.get('APPLY', '').strip().lower() in ('1', 'true', 'yes')
SHEET_URL = ('https://docs.google.com/spreadsheets/d/'
             '1wyP6KnQO5LowvsHZoHM4rntFwyJA8Wc5Kc652SXJDtI/export?format=xlsx')
XLSX = os.environ.get('IMPORT_XLSX', SHEET_URL)
BRANCH = os.environ.get('IMPORT_BRANCH', 'Bahria').strip()
TAB = os.environ.get('IMPORT_TAB', '1st oct 2026').strip()
XMLID_MODULE = '__import__'

COLUMNS = {
    'name': ['name'], 'size': ['size', 'amp/watt/volt'],
    'qty': ['quantity', 'quantities', 'qty', 'pieces'],
    'cost': ['cost'], 'price': ['sales price'], 'min': ['minimum selling price'],
    'mrp': ['maximum retail price (mrp)'], 'categ': ['product category'], 'unit': ['unit'],
}
# Section headings of this tab -> department, matching the Bahria import's
# categories so the products land in the till sections already there.
SECTION_DEPTS = {
    'sanitary_plumbering': ['Sanitary & Plumbing'], 'sanitary_plumbing': ['Sanitary & Plumbing'],
    'sanitary': ['Sanitary & Plumbing'],
    'paint_chemicals': ['Paints', 'Polish', 'Chemicals'], 'chemicals': ['Paints', 'Polish', 'Chemicals'],
    'paints': ['Paints'], 'paint': ['Paints'], 'polish': ['Paints', 'Polish'],
    'electric': ['Electric'], 'hardware': ['Hardware'],
}
CATEGORY_WORDS = {
    'distamber': 'Distemper', 'oilpant': 'Oil Paint', 'oilpaint': 'Oil Paint',
    'electric_items': 'Electric Items', 'gasitems': 'Gas Items',
}
UNIT_WORDS = {
    'ft': 'ft', 'fts': 'ft', 'feet': 'ft', 'foot': 'ft',
    'm': 'm', 'meter': 'm', 'meters': 'm', 'metre': 'm', 'metres': 'm', 'mtr': 'm', 'mtrs': 'm',
    'kg': 'kg', 'kgs': 'kg', 'kilo': 'kg', 'kilos': 'kg',
    'l': 'L', 'lt': 'L', 'ltr': 'L', 'ltrs': 'L', 'lts': 'L', 'litr': 'L',
    'liter': 'L', 'liters': 'L', 'litre': 'L', 'litres': 'L',
}
GRAM_WORDS = {'g', 'gm', 'gms', 'gram', 'grams', 'grm'}
SIZE_UNITS = {'x', 'mm', 'cm', 'ml', 'kg', 'gm', 'ft', 'ltr'}
SECTION_NAMES = {'handbrush': 'Brushes', 'spraypaints': 'Spray Paint', 'spray paints': 'Spray Paint',
                 'paint tube': 'Paint Tubes', 'regmal': 'Sandpaper'}
DEEP_SECTIONS = {'chemicals', 'brackets'}
SECTION_ORDER = ['Distemper', 'Oil Paint', 'Putty', 'Brushes', 'Spray Paint', 'Paint Tubes', 'Sandpaper',
                 'Polish', 'Chemicals', 'Electric Items', 'Brackets']


def clean(value):
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return f"{value.month}/{value.day}"  # Sheets turned a size like 3/4 into a date
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\s+', ' ', str(value)).strip()


def slug(*parts):
    return re.sub(r'[^a-z0-9]+', '_', ' '.join(str(p) for p in parts).lower()).strip('_')


def row_ranges(numbers):
    runs = []
    for n in sorted(numbers):
        if runs and n == runs[-1][1] + 1:
            runs[-1][1] = n
        else:
            runs.append([n, n])
    return ', '.join(f"{a}-{b}" if b > a else f"{a}" for a, b in runs)


def amount(value):
    """480 -> (480, ''), '50 ft' -> (50, 'ft'), '12 pcs' -> (12, ''), empty -> (None, '')."""
    if isinstance(value, (int, float)):
        return float(value), ''
    s = clean(value)
    if not s or s.lower() in ('none', '-', '.'):
        return None, ''
    s = re.sub(r'(\d)[oO](?=\D|$)', r'\g<1>0', s)
    match = re.search(r'\d+(?:\.\d+)?', s)
    if not match:
        return None, ''
    num = float(match.group(0))
    if s.lstrip().startswith('-'):
        num = -num
    unit = next((UNIT_WORDS[w] for w in re.findall(r'[a-z]+', s.lower()) if w in UNIT_WORDS), '')
    return num, unit


def measured_amount(raw):
    """'45Litr' -> (45, 'L'), '2 KILO 850 GRAM' -> (2.85, 'kg'); else (None, '')."""
    s = clean(raw).lower()
    pattern = r'(\d+(?:\.\d+)?)\s*([a-z]+)'
    parts = re.findall(pattern, s)
    if not parts or re.sub(pattern, '', s).strip():
        return None, ''
    total, unit = 0.0, ''
    for number, word in parts:
        if word in GRAM_WORDS:
            part_unit, value = 'kg', float(number) / 1000
        elif word in UNIT_WORDS:
            part_unit, value = UNIT_WORDS[word], float(number)
        else:
            return None, ''
        if unit and part_unit != unit:
            return None, ''
        unit, total = part_unit, total + value
    return round(total, 3), unit


def tidy_size(raw):
    return ' '.join(w if re.search(r'\d', w) or w.lower() in SIZE_UNITS
                    else w.capitalize() if w.isalpha() and (w.islower() or w.isupper())
                    else w for w in raw.split(' '))


def category_path(raw):
    return [CATEGORY_WORDS.get(p.strip().lower(), p.strip().title()) for p in str(raw).split('/') if p.strip()]


def till_section(categ):
    if len(categ) >= 3 and categ[2].lower() in DEEP_SECTIONS:
        sub = categ[2]
    elif len(categ) >= 2:
        sub = categ[1]
    else:
        return categ[0], None
    return categ[0], SECTION_NAMES.get(sub.lower(), sub)


def section_label(categ):
    top, sub = till_section(categ)
    return f"{top} > {sub}" if sub else top


# ------------------------------------------------------------------ read tab
if XLSX.startswith('http'):
    with urllib.request.urlopen(XLSX, timeout=60) as response:
        payload = response.read()
    if not payload.startswith(b'PK'):
        raise SystemExit("The sheet did not download as a spreadsheet -- is it still shared "
                         "'Anyone with the link can view'?")
    source = io.BytesIO(payload)
else:
    source = XLSX
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
tab = next((t for t in wb.sheetnames if t.strip().lower() == TAB.lower()), None)
if not tab:
    raise SystemExit(f"No tab named {TAB!r}. The tabs are: {wb.sheetnames}")
PREFIX = f"tab_{slug(tab)}_"

branch = env['res.company'].sudo().search([('name', 'ilike', BRANCH), ('child_ids', '=', False)])
if len(branch) != 1:
    raise SystemExit(f"Expected exactly one branch matching {BRANCH!r}, found "
                     f"{branch.mapped('name') or 'none'}. Pass the right one as IMPORT_BRANCH.")
warehouse = env['stock.warehouse'].sudo().search([('company_id', '=', branch.id)], limit=1)
if not warehouse:
    raise SystemExit(f"{branch.name} has no warehouse to hold the opening stock.")

Template = env['product.template'].sudo().with_context(active_test=False)
IMD = env['ir.model.data'].sudo()
companies = env['res.company'].sudo().search([])

print()
print("=" * 78)
print(f"IMPORT TAB {tab!r} " + ("" if APPLY else "-- CHECK (nothing is written -- APPLY=1 to update)"))
print("=" * 78)
print(f"branch {branch.name} | warehouse {warehouse.name} | only this tab; the Bahria tabs are not touched")

rows = list(wb[tab].iter_rows(values_only=True))
if not rows:
    raise SystemExit(f"Tab {tab!r} is empty.")
header = [clean(h).lower() for h in rows[0]]
col = {key: next((header.index(a) for a in aliases if a in header), None) for key, aliases in COLUMNS.items()}
brand_col = next((i for i, h in enumerate(header) if h == 'brand'), None)
missing = [k for k in ('name', 'price', 'cost', 'qty') if col[k] is None]
if missing:
    raise SystemExit(f"Tab {tab!r} has no column for {missing}. Header: {[clean(h) for h in rows[0] if h]}")

products, notes, skips, no_price = [], [], [], []
by_key, by_name = {}, {}
dept, blank_before = None, True
for rownum, row in enumerate(rows[1:], start=2):
    get = lambda k: row[col[k]] if col[k] is not None and col[k] < len(row) else None
    if not any(c is not None and str(c).strip() for c in row):
        blank_before = True
        continue
    name = clean(get('name'))
    was_blank, blank_before = blank_before, False
    if not name:
        continue
    size_raw = clean(get('size'))
    size_raw = '' if size_raw in ('0', '-') else size_raw
    brand_raw = clean(row[brand_col]) if brand_col is not None and brand_col < len(row) else ''
    price, p_unit = amount(get('price'))
    cost, c_unit = amount(get('cost'))
    minimum, _ = amount(get('min'))
    mrp, _ = amount(get('mrp'))
    qty, q_unit = amount(get('qty'))
    if qty is None and (p_unit or c_unit):
        in_hand, in_hand_unit = measured_amount(size_raw)
        if in_hand is not None and in_hand_unit == (p_unit or c_unit):
            notes.append(f"row {rownum} {name}: SIZE '{size_raw}' taken as {in_hand:g} {in_hand_unit} in hand")
            qty, q_unit, size_raw = in_hand, in_hand_unit, ''
    categ_raw = clean(get('categ'))
    if not any(v is not None for v in (price, cost, qty, minimum, mrp)) and not size_raw and not categ_raw:
        if was_blank:
            dept = SECTION_DEPTS.get(slug(name), [name.title()])
            skips.append(('section heading', rownum, f"{name} -> {' / '.join(dept)}"))
        else:
            skips.append(('no price, quantity or size', rownum, name))
        continue
    categ = category_path(categ_raw) if categ_raw else list(dept or [tab.title()])
    size = tidy_size(size_raw)
    # As typed ("PPR", "I.I.L"), only an all-lowercase one capitalised ("local").
    brand = '' if not brand_raw or brand_raw.lower() in ('none', '-') else (
        brand_raw.title() if brand_raw.islower() else brand_raw)
    price, cost, minimum, mrp = price or 0.0, cost or 0.0, minimum or 0.0, mrp or 0.0
    if price <= 0:
        no_price.append(rownum)
    elif cost > price:
        notes.append(f"row {rownum} {name}: cost {cost:g} above sales price {price:g}")
    if price and minimum > price:
        notes.append(f"row {rownum} {name}: min {minimum:g} above price {price:g} -> min left out")
        minimum = 0.0
    if price and mrp and mrp < price:
        notes.append(f"row {rownum} {name}: MRP {mrp:g} below price {price:g} -> MRP left out")
        mrp = 0.0
    if minimum and mrp and minimum > mrp:
        minimum = mrp = 0.0
    if qty is not None and qty < 0:
        notes.append(f"row {rownum} {name}: negative quantity {qty:g} ignored")
        qty = None
    units = {u for u in (UNIT_WORDS.get(clean(get('unit')).lower(), ''), p_unit, c_unit, q_unit) if u}
    if len(units) > 1:
        notes.append(f"row {rownum} {name}: mixes units {sorted(units)} -> sold per piece")
    p = {'row': rownum, 'name': name, 'size': size, 'brand': brand, 'price': price, 'cost': cost,
         'minimum': minimum, 'mrp': mrp, 'qty': qty, 'categ': categ,
         'unit': units.pop() if len(units) == 1 else '',
         'key': PREFIX + slug(name, size, brand)}
    first = by_key.get(p['key'])
    if first:
        if qty:
            first['qty'] = (first['qty'] or 0.0) + qty
        notes.append(f"row {rownum} {name}: same name/size/brand as row {first['row']}, merged (quantities added)")
        continue
    twin = by_name.get(slug(name))
    if twin and abs(twin['price'] - price) < 0.01 and (twin['qty'] or 0) == (qty or 0):
        if p['unit'] and not twin['unit']:
            twin['unit'] = p['unit']  # "150 kg" on the copy: the product is sold by the kilo
        skips.append(('repeats an earlier row (same name, price and quantity)', rownum, f"{name} = row {twin['row']}"))
        continue
    by_key[p['key']] = p
    by_name.setdefault(slug(name), p)
    products.append(p)

# ------------------------------------------------- this tab's own products only
mine = {imd.name: imd.res_id for imd in IMD.search([
    ('module', '=', XMLID_MODULE), ('model', '=', 'product.template'), ('name', '=like', 'tab_%')])
    if imd.name.startswith(PREFIX)}  # "_" is a LIKE wildcard: match the prefix exactly here
for p in products:
    p['template'] = Template.browse(mine.get(p['key'])).exists()
claimed = {p['template'].id for p in products if p['template']}


def counted(templates):
    """Templates already counted in the branch (an inventory adjustment exists)."""
    variants = templates.product_variant_ids
    if not variants:
        return set()
    groups = env['stock.move'].sudo()._read_group(
        [('product_id', 'in', variants.ids), ('company_id', '=', branch.id),
         ('is_inventory', '=', True), ('state', '=', 'done')], ['product_id'])
    return {product.product_tmpl_id.id for (product,) in groups}


def untouched(t):
    variants = t.product_variant_ids.ids
    return not any(model in env and env[model].sudo().search_count([('product_id', 'in', variants)], limit=1)
                   for model in ('stock.move', 'pos.order.line', 'sale.order.line', 'purchase.order.line'))


def differences(p):
    t = p['template']
    diff = {}
    compare = [
        ('name', t.name or '', p['name']), ('size', t.pos_retail_size or '', p['size']),
        ('brand', t.brand_id.name or '', p['brand']), ('price', t.list_price, p['price']),
        ('min', t.minimum_selling_price, p['minimum']), ('MRP', t.mrp, p['mrp']),
        ('category', t.categ_id.complete_name or '', ' / '.join(p['categ'])),
        ('till section', ' > '.join(n for n in (t.pos_categ_ids[:1].parent_id.name, t.pos_categ_ids[:1].name) if n)
         if t.pos_categ_ids else '', section_label(p['categ'])),
    ]
    for field, now, want in compare:
        if isinstance(want, float):
            same = abs((now or 0.0) - want) <= 0.001
        elif field == 'name':
            same = clean(now) == clean(want)
        elif field in ('size', 'brand'):
            same = slug(now) == slug(want)
        else:
            same = now == want
        if not same:
            diff[field] = (f"{now:g}", f"{want:g}") if isinstance(want, float) else (now or '-', want or '-')
    if any(abs(t.with_company(c).standard_price - p['cost']) > 0.001 for c in companies):
        diff['cost'] = (f"{t.with_company(branch).standard_price:g}", f"{p['cost']:g}")
    if p['unit'] and t.uom_id.name != p['unit'] and untouched(t):
        diff['unit'] = (t.uom_id.name, p['unit'])
    if t.pos_retail_branch_ids != branch:
        diff['sold at'] = (', '.join(t.pos_retail_branch_ids.mapped('name')) or 'every branch', branch.name)
    if not (t.available_in_pos and t.sale_ok and t.purchase_ok and t.is_storable and not t.company_id):
        diff['settings'] = ('off', 'sellable, stocked, shared')
    return diff


# ------------------------------------------------------------------- report
already = counted(Template.browse(list(claimed)))
new = [p for p in products if not p['template']]
changed = [(p, d) for p in products if p['template'] for d in [differences(p)] if d]
stock_load = [p for p in products if p['qty'] and (not p['template'] or p['template'].id not in already)]
print(f"\n{tab}: {len(products)} products  |  {len(new)} new, {len(changed)} to update, "
      f"{len(products) - len(new) - len(changed)} unchanged")
print(f"  categories:    {sorted({' / '.join(p['categ']) for p in products})}")
print(f"  till sections: {sorted({section_label(p['categ']) for p in products})}")
print(f"  opening stock to load: {len(stock_load)} products"
      + (f" (units: {sorted({p['unit'] for p in products if p['unit']})})" if any(p['unit'] for p in products) else ""))
if new:
    print(f"  NEW ({len(new)}):")
    for p in new[:40]:
        print(f"    + row {p['row']} {p['name']} | {p['size'] or '-'} | {p['brand'] or '-'} | "
              f"{' / '.join(p['categ'])} | price {p['price']:g} | qty "
              f"{p['qty'] if p['qty'] is not None else '-'}{' ' + p['unit'] if p['unit'] else ''}")
    if len(new) > 40:
        print(f"    ... and {len(new) - 40} more")
if changed:
    print(f"  TO UPDATE ({len(changed)}):")
    for p, d in changed[:40]:
        print(f"    ~ row {p['row']} {p['name']}: " + '; '.join(f"{f} {a} -> {b}" for f, (a, b) in d.items()))
empty_qty = [p for p in products if not p['qty']]
if empty_qty:
    print(f"  NO QUANTITY ({len(empty_qty)}) -- will show out of stock: rows {row_ranges([p['row'] for p in empty_qty])}")
if no_price:
    print(f"  NO SALES PRICE ({len(no_price)}) -- the till won't sell these until the sheet has one: "
          f"rows {row_ranges(no_price)}")
for n in notes:
    print(f"  note: {n}")
groups = {}
for reason, rownum, name in skips:
    groups.setdefault(reason, []).append(f"{rownum} {name}")
for reason, items in groups.items():
    print(f"  skipped -- {reason} ({len(items)}): rows {', '.join(items[:8])}"
          + (f" ... +{len(items) - 8}" if len(items) > 8 else ""))
gone = Template.browse([rid for rid in mine.values() if rid not in claimed]).exists().filtered('active')
if gone:
    print(f"\n{len(gone)} product(s) this script made earlier are no longer in the tab (left as they are): "
          f"{', '.join(gone.mapped('name')[:20])}")

if not APPLY:
    print("\n" + "=" * 78)
    print("CHECK ONLY -- nothing was written. Run again with APPLY=1 to import.")
    print("=" * 78)
else:
    Category = env['product.category'].sudo()
    PosCategory = env['pos.category'].sudo()
    Brand = env['product.brand'].sudo()
    caches = {'categ': {}, 'section': {}, 'brand': {}, 'uom': {}}

    def category(path):
        key = tuple(path)
        if key not in caches['categ']:
            parent = Category.browse()
            for name in path:
                parent = Category.search([('name', '=', name), ('parent_id', '=', parent.id or False)], limit=1) \
                    or Category.create({'name': name, 'parent_id': parent.id or False})
            caches['categ'][key] = parent
        return caches['categ'][key]

    def section(categ):
        top, sub = till_section(categ)
        if (top, sub) not in caches['section']:
            parent = PosCategory.search([('name', '=ilike', top), ('parent_id', '=', False)], limit=1) \
                or PosCategory.create({'name': top})
            record = parent
            if sub:
                record = PosCategory.search([('name', '=ilike', sub), ('parent_id', '=', parent.id)], limit=1) \
                    or PosCategory.create({'name': sub, 'parent_id': parent.id,
                                           'sequence': SECTION_ORDER.index(sub) if sub in SECTION_ORDER else 99})
            caches['section'][(top, sub)] = record
        return caches['section'][(top, sub)]

    def brand(name):
        if not name:
            return False
        if name.lower() not in caches['brand']:
            caches['brand'][name.lower()] = Brand.with_context(active_test=False).search(
                [('name', '=ilike', name)], limit=1) or Brand.create({'name': name})
        return caches['brand'][name.lower()].id

    def uom(unit):
        if unit not in caches['uom']:
            found = env['uom.uom'].sudo().search([('name', '=', unit)], limit=1) if unit else None
            caches['uom'][unit] = found or env.ref('uom.product_uom_unit')
        return caches['uom'][unit]

    def full_vals(p):
        return {
            'name': p['name'], 'pos_retail_size': p['size'] or False, 'brand_id': brand(p['brand']),
            'list_price': p['price'], 'minimum_selling_price': p['minimum'], 'mrp': p['mrp'],
            'categ_id': category(p['categ']).id, 'pos_categ_ids': [(6, 0, section(p['categ']).ids)],
            'uom_id': uom(p['unit']).id,
            'pos_retail_branch_ids': [(6, 0, branch.ids)], 'type': 'consu', 'is_storable': True,
            'available_in_pos': True, 'sale_ok': True, 'purchase_ok': True, 'company_id': False,
        }

    FIELD_VALS = {
        'name': ['name'], 'size': ['pos_retail_size'], 'brand': ['brand_id'], 'price': ['list_price'],
        'min': ['minimum_selling_price'], 'MRP': ['mrp'], 'category': ['categ_id'],
        'till section': ['pos_categ_ids'], 'sold at': ['pos_retail_branch_ids'], 'unit': ['uom_id'],
        'settings': ['type', 'is_storable', 'available_in_pos', 'sale_ok', 'purchase_ok', 'company_id'],
    }
    print("\nimporting...", flush=True)
    with env.cr.savepoint():
        updated = 0
        for p, d in changed:
            vals = full_vals(p)
            write = {k: vals[k] for f in d for k in FIELD_VALS.get(f, []) if k in vals}
            if write:
                p['template'].write(write)
                updated += 1
        if new:
            records = Template.create([full_vals(p) for p in new])
            IMD.create([{'module': XMLID_MODULE, 'name': p['key'], 'model': 'product.template',
                         'res_id': record.id, 'noupdate': True} for p, record in zip(new, records)])
            for p, record in zip(new, records):
                p['template'] = record
        for company in companies:
            by_cost = {}
            for p in products:
                if abs(p['template'].with_company(company).standard_price - p['cost']) > 0.001:
                    by_cost.setdefault(p['cost'], Template.browse())
                    by_cost[p['cost']] |= p['template']
            for value, records in by_cost.items():
                records.with_company(company).write({'standard_price': value})
        wanted = {p['template']: p['qty'] for p in products if p['qty']}
        done = counted(Template.browse([t.id for t in wanted]))
        load = {t: q for t, q in wanted.items() if t.id not in done}
        if load:
            quants = env['stock.quant'].sudo().with_company(branch).with_context(inventory_mode=True).create([{
                'product_id': t.product_variant_id.id, 'location_id': warehouse.lot_stock_id.id,
                'inventory_quantity': q} for t, q in load.items()])
            if isinstance(quants.action_apply_inventory(), dict):
                raise SystemExit("Odoo asked to resolve an inventory conflict; nothing was saved. "
                                 "Check Inventory > Physical Inventory for this branch.")
    env.cr.commit()
    print(f"  {tab}: created {len(new)}, updated {updated}, opening stock on {len(load)} products")
    print("=" * 78)
    print("Import finished. Reload the till pages.")
    print("=" * 78)
