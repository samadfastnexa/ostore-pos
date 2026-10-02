"""Check and import the whole Bahria catalogue from the Murshid Store Google Sheet.

    Server:  bash scripts/run_import_on_server.sh           (check -- writes nothing)
             APPLY=1 bash scripts/run_import_on_server.sh   (update products)
    Local:   IMPORT_BRANCH="Branch 2" venv/Scripts/python.exe odoo/odoo-bin shell \
                 -c odoo.conf -d OStore --no-http < custom_addons/pos_retail/scripts/import_catalogue.py

Every tab whose name starts with "Bahria" is read -- a new tab is picked up
without touching this file. Columns are found by header name, whatever order
they are in.

WITHOUT APPLY=1 THIS IS A CHECK: it compares the sheet with Odoo and prints,
per tab, the new products, every field that would change on existing ones
("price 400 -> 450"), stock that would be loaded, products that will show
out of stock and why, and the rows it skips. Nothing is written.

Environment (all optional):
    APPLY=1          write the changes
    IMPORT_BRANCH    branch that sells and stocks the products (default Bahria)
    IMPORT_TABS      only these tabs, comma-separated
    IMPORT_XLSX      a local .xlsx instead of the Google Sheet
    FORCE_STOCK=1    re-count products already counted -- OVERWRITES live stock

HOW A ROW BECOMES A PRODUCT
  * identity: the BARCODE code (MBAHRIA0001...) when there is one -- it is set
    as the product's barcode and Internal Reference. A row without a code is
    still imported, recognised on later runs by tab + name + size + brand, and
    gets an automatic barcode until a code is added to the sheet (the next run
    then moves the product onto that code). Codes must never be renumbered:
    the run refuses when many codes suddenly sit on differently named products.
  * a corrected row stays the same product: when a row no longer finds its
    product (brand or size fixed, the two columns swapped back, code changed,
    the name made longer: "Tee" -> "Tee - 11L") it keeps the one unmatched
    product of the same department that plainly fits -- listed as "[same
    product, row changed]". A second copy an earlier run made that way
    (nothing but its opening count) is archived with that stock zeroed; the
    original keeps its stock and sales.
  * counted in the wrong unit (15 pipes in Units, now "195ft" per foot) and
    never sold: replaced by a product in the sheet's unit, counted afresh.
  * gone from the sheet and never sold or bought: archived, stock zeroed.
    With sales it is only listed. Never a quarter or more of a department at
    once (a renamed tab or a bad download), never a tab not imported now.
  * a heading row named NEW (or "new items"...) is a marker, not a category.
  * name exactly as in the sheet; size into the Size field; brand from the
    first Brand column with a value; sold at the branch only; cost set in
    every company (company-dependent in Odoo 19).
  * category: the Product Category column; else the heading row above it (a
    row with only a name, after an empty row, e.g. "CONDUTE eLECTRIC"); else
    the tab's department. Till section = the category's second level.
  * a price, cost or quantity with a unit ("160 rs/ft", "480/lts", "242 total
    fts", "6 Kg") creates the product in that unit (ft, L, kg, m); the till
    then sells it in fractions (1.5 L at 300/L = 450). An existing product is
    moved to that unit only while it has no stock, sales or purchases yet.
  * loose goods sold per unit with QUANTITY empty: a SIZE in that unit is the
    amount in hand ("Thinner | 17Litr | | 470/ltr" -> 17 L; "2 KILO 850 GRAM"
    -> 2.85 kg), and the size itself is left blank.
  * opening stock = the sheet's quantity, loaded ONCE per product: anything
    already counted in the branch (by this import or by hand) is never
    touched again -- sales and receipts keep it right after that.
  * skipped and listed, never imported: rows pasted far below the table (after
    10+ empty rows) without a code, copies of a product that has a code
    elsewhere in the sheet, heading rows, and rows with no price, quantity or
    code at all.
"""

import datetime
import io
import os
import re
import time
import urllib.request

import openpyxl

from odoo.exceptions import UserError

APPLY = os.environ.get('APPLY', '').strip().lower() in ('1', 'true', 'yes')
FORCE_STOCK = os.environ.get('FORCE_STOCK', '').strip().lower() in ('1', 'true', 'yes')
SHEET_URL = ('https://docs.google.com/spreadsheets/d/'
             '1wyP6KnQO5LowvsHZoHM4rntFwyJA8Wc5Kc652SXJDtI/export?format=xlsx')
XLSX = os.environ.get('IMPORT_XLSX', SHEET_URL)
BRANCH = os.environ.get('IMPORT_BRANCH', 'Bahria').strip()
ONLY_TABS = [t.strip().lower() for t in os.environ.get('IMPORT_TABS', '').split(',') if t.strip()]
TAB_PREFIX = 'bahria'
GAP = 10
# A heading row with one of these names marks new stock, it is not a category.
MARKER_HEADINGS = {'new', 'new items', 'new stock', 'new arrival', 'new arrivals'}
XMLID_MODULE = '__import__'
IMPORT_PREFIXES = ('cat_', 'paint_', 'polish_', 'chemical_', 'electric_')

# Department and default category per tab (anything else: the tab's own name).
DEFAULT_CATEG = {
    'bahria paints': ['Paints'], 'bahria polish': ['Paints', 'Polish'],
    'bahria chemicals': ['Paints', 'Polish', 'Chemicals'], 'bahria electric': ['Electric'],
    'bahria hardware': ['Hardware'], 'bahria sanitoryplumbering': ['Sanitary & Plumbing'],
}
COLUMNS = {
    'name': ['name'], 'size': ['size', 'amp/watt/volt'], 'qty': ['quantity', 'pieces'],
    'cost': ['cost'], 'price': ['sales price'], 'min': ['minimum selling price'],
    'mrp': ['maximum retail price (mrp)'], 'categ': ['product category'],
    'code': ['barcode'], 'unit': ['unit'],
}

# Spellings below are what earlier imports produced -- change one and the
# products it covers are renamed/re-filed on the next run. Only ADD entries.
SIZES = {
    'gallon': 'Gallon', 'quarter': 'Quarter', 'drum': 'Drum',
    'adha pound': 'Half Pound', 'half liter': 'Half Litre',
    '4 inche': '4 Inch', '4 inchi': '4 Inch', '5 inchi': '5 Inch',
}
BRANDS = {
    'advacne': 'Advance', 'advance series': 'Advance',
    'black and white': 'Black and White', 'captain': 'Captain', 'commander': 'Commander',
    'elite': 'Elite', 'exclusive nelson': 'Exclusive Nelson', 'nelson exclusive': 'Exclusive Nelson',
    'fine coat': 'Finecoat', 'finecoat': 'Finecoat',
    'finecoat/ makro': 'Makro / Finecoat', 'makro/fincoat': 'Makro / Finecoat',
    'marko/fincoat': 'Makro / Finecoat', 'fish': 'Fish', 'glide': 'Glide',
    "gobi's": "Gobi's", 'group master': 'Group Master', 'jotun': 'Jotun',
    'kent tone': 'Kent Tone', 'local': 'Local', 'makro': 'Makro',
    'murshid colors': 'Murshid Colors', 'nelson': 'Nelson', 'nelson extra': 'Nelson Extra',
    'silicon master paint': 'Silicon Master Paint', 'sooper': 'Sooper', 'universal': 'Universal',
    'aqua': 'AQUA', 'arik-lux': 'Arik-Lux', 'brooks': 'Brooks', 'broox': 'Broox',
    'burq': 'Burq', 'china': 'China', 'classic': 'Classic', 'clopal': 'Clopal',
    'chinton': 'Chinton', 'daiichi': 'Daiichi', 'deco': 'Deco', 'diamond': 'Diamond',
    'excelent': 'Excellent', 'five star': 'Five Star', 'flud': 'Flud',
    'meezan': 'Meezan', 'max': 'MAX', 'schneider': 'Schneider',
}
CATEGORY_WORDS = {
    'distamber': 'Distemper', 'oilpant': 'Oil Paint', 'oilpaint': 'Oil Paint',
    'electric_items': 'Electric Items', 'gasitems': 'Gas Items',
}
UNIT_WORDS = {
    'ft': 'ft', 'fts': 'ft', 'feet': 'ft', 'foot': 'ft',
    'm': 'm', 'meter': 'm', 'meters': 'm', 'metre': 'm', 'metres': 'm', 'mtr': 'm', 'mtrs': 'm',
    'kg': 'kg', 'kgs': 'kg', 'kilo': 'kg', 'kilos': 'kg',
    'l': 'L', 'lt': 'L', 'ltr': 'L', 'ltrs': 'L', 'lts': 'L', 'litr': 'L',
    'liter': 'L', 'liters': 'L', 'litre': 'L', 'litres': 'L',
}
# Till sections: the second category level, renamed where the category word
# is not what a cashier would look for.
SECTION_NAMES = {'handbrush': 'Brushes', 'spraypaints': 'Spray Paint', 'spray paints': 'Spray Paint',
                 'paint tube': 'Paint Tubes', 'regmal': 'Sandpaper'}
DEEP_SECTIONS = {'chemicals', 'brackets'}
SECTION_ORDER = ['Distemper', 'Oil Paint', 'Putty', 'Brushes', 'Spray Paint', 'Paint Tubes', 'Sandpaper',
                 'Polish', 'Chemicals', 'Electric Items', 'Brackets']


def clean(value):
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return f"{value.month}/{value.day}"  # Sheets turned a size like 1/2 into a date
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\s+', ' ', str(value)).strip()


def slug(*parts):
    return re.sub(r'[^a-z0-9]+', '_', ' '.join(str(p) for p in parts).lower()).strip('_')


def row_ranges(numbers):
    """[6, 7, 42, 43, 44, 50] -> '6-7, 42-44, 50'"""
    runs = []
    for n in sorted(numbers):
        if runs and n == runs[-1][1] + 1:
            runs[-1][1] = n
        else:
            runs.append([n, n])
    return ', '.join(f"{a}-{b}" if b > a else f"{a}" for a, b in runs)


def amount(value):
    """480 -> (480, ''), '480/lts' -> (480, 'L'), '242 total fts' -> (242, 'ft'),
    '7500.kg' -> (7500, 'kg'), '35o' -> (350, ''), empty -> (None, '')."""
    if isinstance(value, (int, float)):
        return float(value), ''
    s = clean(value)
    if not s or s.lower() in ('none', '-', '.'):
        return None, ''
    s = re.sub(r'(\d)[oO](?=\D|$)', r'\g<1>0', s)
    match = re.search(r'\d+(?:\.\d+)?', s)
    if not match:
        return None, ''
    num = float(match.group(0))
    if s.lstrip().startswith('-'):
        num = -num
    unit = next((UNIT_WORDS[w] for w in re.findall(r'[a-z]+', s.lower()) if w in UNIT_WORDS), '')
    return num, unit


GRAM_WORDS = {'g', 'gm', 'gms', 'gram', 'grams', 'grm'}


def measured_amount(raw):
    """A size that is really an amount: '45Litr' -> (45, 'L'), '2 KILO 850 GRAM'
    -> (2.85, 'kg'), '1170 gram' -> (1.17, 'kg'); anything else -> (None, '')."""
    s = clean(raw).lower()
    pattern = r'(\d+(?:\.\d+)?)\s*([a-z]+)'
    parts = re.findall(pattern, s)
    if not parts or re.sub(pattern, '', s).strip():
        return None, ''
    total, unit = 0.0, ''
    for number, word in parts:
        if word in GRAM_WORDS:
            part_unit, value = 'kg', float(number) / 1000
        elif word in UNIT_WORDS:
            part_unit, value = UNIT_WORDS[word], float(number)
        else:
            return None, ''
        if unit and part_unit != unit:
            return None, ''
        unit, total = part_unit, total + value
    return round(total, 3), unit


SIZE_UNITS = {'x', 'mm', 'cm', 'ml', 'kg', 'gm', 'ft', 'ltr'}  # '32 x 25 mm', not '32 X 25 Mm'


def tidy_size(raw):
    """Known sizes to their usual spelling; otherwise capitalise plain words and
    leave anything with a digit alone ('25MM', '3/4"', '1liter' stay as typed)."""
    if raw.lower() in SIZES:
        return SIZES[raw.lower()]
    return ' '.join(w if re.search(r'\d', w) or w.lower() in SIZE_UNITS
                    else w.capitalize() if w.isalpha() and (w.islower() or w.isupper())
                    else w for w in raw.split(' '))


def category_path(raw):
    return [CATEGORY_WORDS.get(p.strip().lower(), p.strip().title()) for p in str(raw).split('/') if p.strip()]


def till_section(categ):
    """('Paints', 'Distemper') for Paints/Distemper/Drum; (top, None) for a bare top category."""
    if len(categ) >= 3 and categ[2].lower() in DEEP_SECTIONS:
        sub = categ[2]
    elif len(categ) >= 2:
        sub = categ[1]
    else:
        return categ[0], None
    return categ[0], SECTION_NAMES.get(sub.lower(), sub)


# ---------------------------------------------------------------- read sheet
if XLSX.startswith('http'):
    with urllib.request.urlopen(XLSX, timeout=120) as response:
        payload = response.read()
    if not payload.startswith(b'PK'):
        raise SystemExit("The sheet did not download as a spreadsheet -- is it still shared "
                         "'Anyone with the link can view'?")
    source = io.BytesIO(payload)
else:
    source = XLSX
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
# Every Bahria tab is read and matched, so a product on another tab is never
# taken for a stray; IMPORT_TABS only narrows what is reported and written.
tabs = [t for t in wb.sheetnames if t.lower().startswith(TAB_PREFIX)]
work_tabs = [t for t in tabs if not ONLY_TABS or t.lower() in ONLY_TABS or any(o in t.lower() for o in ONLY_TABS)]
if not work_tabs:
    raise SystemExit(f"No tab matches IMPORT_TABS {ONLY_TABS}; the Bahria tabs are {tabs}.")

branch = env['res.company'].sudo().search([('name', 'ilike', BRANCH), ('child_ids', '=', False)])
if len(branch) != 1:
    raise SystemExit(f"Expected exactly one branch matching {BRANCH!r}, found "
                     f"{branch.mapped('name') or 'none'}. Pass the right one as IMPORT_BRANCH.")
warehouse = env['stock.warehouse'].sudo().search([('company_id', '=', branch.id)], limit=1)
if not warehouse:
    raise SystemExit(f"{branch.name} has no warehouse to hold the opening stock.")

Template = env['product.template'].sudo().with_context(active_test=False)
IMD = env['ir.model.data'].sudo()
companies = env['res.company'].sudo().search([])

print()
print("=" * 78)
print("MURSHID CATALOGUE " + ("IMPORT" if APPLY else "CHECK   (nothing is written -- APPLY=1 to update)"))
print("=" * 78)
print(f"branch {branch.name} | warehouse {warehouse.name} | tabs {work_tabs}")

# ------------------------------------------------------------- parse all tabs
parsed = {}          # tab -> list of product dicts
skipped = {}         # tab -> list of (reason, row, name)
tab_notes = {}
tab_no_price = {}    # tab -> row numbers without a sales price
for tab in tabs:
    rows = list(wb[tab].iter_rows(values_only=True))
    if not rows:
        continue
    header = [clean(h).lower() for h in rows[0]]
    col = {key: next((header.index(a) for a in aliases if a in header), None) for key, aliases in COLUMNS.items()}
    # Only the first Brand column: in Hardware and Sanitary a second one holds
    # leftovers of pasted paint rows ("Exclusive Nelson" on a gas part).
    brand_cols = [i for i, h in enumerate(header) if h == 'brand'][:1]
    missing = [k for k in ('name', 'price', 'cost', 'qty') if col[k] is None]
    if missing:
        raise SystemExit(f"Tab {tab!r} has no column for {missing}. Header: "
                         f"{[clean(h) for h in rows[0] if h]}. Rename it back.")
    default = DEFAULT_CATEG.get(tab.lower(), [re.sub(r'(?i)^bahria\s*', '', tab).strip().title() or tab])
    products, notes, skips, no_price = [], [], [], []
    heading, empty_run, blank_before, in_table = None, 0, True, True
    for rownum, row in enumerate(rows[1:], start=2):
        get = lambda k: row[col[k]] if col[k] is not None and col[k] < len(row) else None
        if not any(c is not None and str(c).strip() for c in row):
            empty_run += 1
            blank_before = True
            if empty_run >= GAP:
                in_table = False
            continue
        name_cell = get('name')
        name = clean(name_cell)
        was_blank, blank_before, empty_run = blank_before, False, 0
        if not name:
            continue
        code = clean(get('code')).upper()
        code = code if re.fullmatch(r'MBAHRIA\d+', code) else ''
        size_raw = clean(get('size'))
        size_raw = '' if size_raw in ('0', '-') else size_raw
        brand_raw = next((clean(row[i]) for i in brand_cols if i < len(row) and clean(row[i])), '')
        price, p_unit = amount(get('price'))
        cost, c_unit = amount(get('cost'))
        minimum, _ = amount(get('min'))
        mrp, _ = amount(get('mrp'))
        qty, q_unit = amount(get('qty'))
        # Loose goods sold by the litre or kilo carry the amount in hand in the
        # SIZE column ("Thinner | 17Litr | (no quantity) | 470/ltr"): a drum's
        # size is meaningless when the price is per litre, its contents are not.
        sold_per = p_unit or c_unit
        if qty is None and sold_per:
            in_hand, in_hand_unit = measured_amount(size_raw)
            if in_hand is not None and in_hand_unit == sold_per:
                notes.append(f"row {rownum} {name}: sold per {sold_per} with QUANTITY empty -> SIZE "
                             f"'{size_raw}' taken as {in_hand:g} {sold_per} in hand")
                qty, q_unit, size_raw = in_hand, in_hand_unit, ''
        categ_raw = clean(get('categ'))
        has_data = any(v is not None for v in (price, cost, qty, minimum, mrp)) or size_raw or categ_raw or code
        if not has_data:
            if was_blank and in_table and name.lower() in MARKER_HEADINGS:
                heading = None
                skips.append(('marker row, not a category', rownum, name))
            elif was_blank and in_table:
                heading = name  # e.g. "CONDUTE eLECTRIC" -- files the rows below it
                skips.append(('heading row (used as the category of the rows below)', rownum, name))
            else:
                skips.append(('no price, quantity or code', rownum, name))
            continue
        if not code and not in_table:
            skips.append(('pasted far below the table, no code', rownum, name))
            continue
        size = tidy_size(size_raw)
        brand = BRANDS.get(brand_raw.lower(), brand_raw.title()) if brand_raw and brand_raw.lower() != 'none' else ''
        if categ_raw:
            categ = category_path(categ_raw)
        elif heading:
            categ = default[:1] + [heading.title()]
        else:
            categ = list(default)
        price = price or 0.0
        cost = cost or 0.0
        minimum = minimum or 0.0
        mrp = mrp or 0.0
        if price <= 0:
            no_price.append(rownum)
        elif cost > price:
            notes.append(f"row {rownum} {name}: cost {cost:g} above sales price {price:g}")
        # A bound that contradicts the price is left out; the other one is kept.
        if price and minimum > price:
            notes.append(f"row {rownum} {name}: min {minimum:g} above price {price:g} -> min left out")
            minimum = 0.0
        if price and mrp and mrp < price:
            notes.append(f"row {rownum} {name}: MRP {mrp:g} below price {price:g} -> MRP left out")
            mrp = 0.0
        if minimum and mrp and minimum > mrp:
            notes.append(f"row {rownum} {name}: min {minimum:g} above MRP {mrp:g} -> range left out")
            minimum = mrp = 0.0
        if qty is not None and qty < 0:
            notes.append(f"row {rownum} {name}: negative quantity {qty:g} ignored")
            qty = None
        units = {u for u in (clean(get('unit')).lower() and UNIT_WORDS.get(clean(get('unit')).lower(), ''),
                             p_unit, c_unit, q_unit) if u}
        if len(units) > 1:
            notes.append(f"row {rownum} {name}: mixes units {sorted(units)} -> sold per piece")
        products.append({
            'tab': tab, 'row': rownum, 'name': name, 'code': code, 'size': size, 'brand': brand,
            'price': price, 'cost': cost, 'minimum': minimum, 'mrp': mrp, 'qty': qty,
            'unit': units.pop() if len(units) == 1 else '', 'categ': categ,
            'key': f"cat_{slug(tab)}_{slug(name, size, brand)}",
        })
    parsed[tab], skipped[tab], tab_notes[tab], tab_no_price[tab] = products, skips, notes, no_price

# Copies: an uncoded row naming the same product (name + size) as a coded row
# anywhere in the sheet is a pasted duplicate, not a second product.
coded_names = {(slug(p['name']), slug(p['size'])) for ps in parsed.values() for p in ps if p['code']}
for tab, products in parsed.items():
    keep, seen = [], {}
    for p in products:
        if not p['code'] and (slug(p['name']), slug(p['size'])) in coded_names:
            skipped[tab].append(('copy of a product with a code elsewhere in the sheet', p['row'], p['name']))
            continue
        ident = p['code'] or p['key']
        if ident in seen:
            first = seen[ident]
            if p['code']:
                raise SystemExit(f"{tab}: code {p['code']} is on rows {first['row']} and {p['row']} -- "
                                 f"every product needs its own code.")
            if p['qty']:
                first['qty'] = (first['qty'] or 0.0) + p['qty']
            tab_notes[tab].append(f"row {p['row']} {p['name']}: same name/size/brand as row {first['row']}, "
                                  f"merged (quantities added)")
            continue
        seen[ident] = p
        keep.append(p)
    parsed[tab] = keep

# ----------------------------------------------------- match to Odoo products
imported_ids = {imd.name: imd.res_id for imd in IMD.search([
    ('module', '=', XMLID_MODULE), ('model', '=', 'product.template')])
    if imd.name.startswith(IMPORT_PREFIXES)}
all_codes = [p['code'] for ps in parsed.values() for p in ps if p['code']]
by_code = {t.default_code.upper(): t for t in Template.search([('default_code', 'in', all_codes)])} if all_codes else {}
claimed = set()
for tab, products in parsed.items():
    renamed = 0
    code_hits = 0
    for p in products:
        t = by_code.get(p['code']) if p['code'] else None
        if t:
            code_hits += 1
            if slug(t.name) != slug(p['name']):
                renamed += 1
        else:
            t = Template.browse(imported_ids.get(p['key'])).exists()
        if t and t.id in claimed:
            t = Template
        p['template'] = t
        if t:
            claimed.add(t.id)
    if code_hits >= 10 and renamed > 0.3 * code_hits:
        raise SystemExit(f"Refusing: in {tab!r}, {renamed} of {code_hits} codes now sit on a differently "
                         f"named product. It looks like the codes were renumbered; put each code back on "
                         f"its own product first.")

# ------------------------------------------- same product, changed in the sheet
# A row whose code or name/size/brand no longer finds its product is usually a
# product already in Odoo whose row was corrected: a brand typo fixed, the size
# and brand columns swapped back, a code added or changed, the name made longer
# ("Tee" -> "Tee - 11L"). Creating it again would put a second copy at the till
# with the stock counted twice. So such a row keeps the unmatched product of the
# same department that is plainly the same thing, when exactly one is:
#
#   same name, same size and brand (either column)        best
#   same name, same size or same brand
#   one name inside the other word for word, same size and brand
#
# and no other row wants that product.


def fit(p, t):
    """2: same size and brand, in either column; 1: one of them; 0: neither."""
    sheet = [slug(p['size']), slug(p['brand'])]
    odoo = [slug(t.pos_retail_size or ''), slug(t.brand_id.name or '')]
    if sorted(sheet) == sorted(odoo):
        return 2
    return 1 if sheet[0] == odoo[0] or sheet[1] == odoo[1] else 0


def words(name):
    return set(slug(name).split('_')) - {''}


def same_product(p, t):
    """How surely the sheet row and the Odoo product are one product (0: not).
    The share of words in common breaks ties: "F.M Wall Socket 11L" is nearer
    "F.M Wall Socket" than "M. Wall Socket", though both names fit inside it."""
    sizes_brands = fit(p, t)
    if slug(p['name']) == slug(t.name):
        return {2: 5, 1: 4}.get(sizes_brands, 0)
    a, b = words(p['name']), words(t.name)
    if a and b and (a <= b or b <= a) and sizes_brands == 2:
        return 2 + len(a & b) / len(a | b)
    return 0


def department(t):
    return (t.categ_id.complete_name or '').split(' / ')[0].lower()


def only_counted(t):
    """Nothing but inventory counts: never sold, bought, quoted or moved."""
    variants = t.product_variant_ids.ids
    if env['stock.move'].sudo().search_count([('product_id', 'in', variants), ('is_inventory', '=', False)], limit=1):
        return False
    return not any(model in env and env[model].sudo().search_count([('product_id', 'in', variants)], limit=1)
                   for model in ('pos.order.line', 'sale.order.line', 'purchase.order.line'))


def untouched(t):
    """Never stocked, sold, bought or quoted -- so its unit can still change.
    (Odoo 19 would relabel the history instead: "10 Units" becoming "10 kg".)"""
    variants = t.product_variant_ids.ids
    return not any(model in env and env[model].sudo().search_count([('product_id', 'in', variants)], limit=1)
                   for model in ('stock.move', 'pos.order.line', 'sale.order.line', 'purchase.order.line'))


imported = set(imported_ids.values())
pool = {}  # department -> imported products no row claims
for t in Template.browse(imported).exists().filtered('active'):
    if t.id not in claimed:
        pool.setdefault(department(t), []).append(t)

# Unmatched rows look for their product; so do rows matched to a copy an
# earlier run made from a renamed row (only an OLDER product -- lower id -- can be the
# original). The copy -- nothing but its opening count -- is then archived with
# that stock zeroed, and the original keeps its stock and sales.
def clear_winner(scored):
    """The best of [(score, item)] when it beats the runner-up, else None."""
    scored = sorted(scored, key=lambda pair: pair[0], reverse=True)
    if scored and (len(scored) == 1 or scored[0][0] > scored[1][0]):
        return scored[0]
    return None


# In rounds: a product one row takes is out of the running for the others, which
# can leave another row with a single clear match.
seeking = [p for ps in parsed.values() for p in ps if not p['template'] or p['template'].id in imported]
while True:
    offers = {}  # product id -> [(score, row)]
    for p in seeking:
        t = p['template']
        scored = []
        for u in pool.get(p['categ'][0].lower(), []):
            # A row that already has its product swaps it only for an older
            # one (lower id) with the very same size and brand.
            if t and (u.id >= t.id or fit(p, u) != 2):
                continue
            score = same_product(p, u)
            if score:
                scored.append((score, u))
        best = clear_winner(scored)
        # Only a copy with nothing but its opening count may be given up.
        if best and (not t or only_counted(t)):
            offers.setdefault(best[1].id, []).append((best[0], p))
    taken = []
    for uid, bids in offers.items():
        winner = clear_winner(bids)
        if not winner:
            continue
        p, t = winner[1], winner[1]['template']
        u = Template.browse(uid)
        if t:
            p['retire'] = [t]
            claimed.discard(t.id)
        else:
            p['kept'] = True
        p['template'] = u
        claimed.add(u.id)
        taken.append((p, u))
    if not taken:
        break
    for p, u in taken:
        seeking.remove(p)
        pool[department(u)].remove(u)

# Counted in the wrong unit: "15" pipes in Units, now "195ft" priced per foot.
# A product with nothing but that count is replaced by one in the sheet's unit,
# counted from the sheet; relabelling would turn 15 pipes into 15 ft.
for products in parsed.values():
    for p in products:
        t = p['template']
        if (t and p['unit'] and t.uom_id.name != p['unit'] and not untouched(t)
                and only_counted(t)):
            p.setdefault('retire', []).append(t)
            p['replaces'] = t
            claimed.discard(t.id)
            p['template'] = Template


def counted(templates):
    """Templates already counted in the branch (an inventory adjustment exists)."""
    variants = templates.product_variant_ids
    if not variants:
        return set()
    groups = env['stock.move'].sudo()._read_group(
        [('product_id', 'in', variants.ids), ('company_id', '=', branch.id),
         ('is_inventory', '=', True), ('state', '=', 'done')], ['product_id'])
    return {product.product_tmpl_id.id for (product,) in groups}


existing = Template.browse([p['template'].id for ps in parsed.values() for p in ps if p['template']])
already_counted = counted(existing)


def section_label(categ):
    top, sub = till_section(categ)
    return f"{top} > {sub}" if sub else top


def differences(p):
    """Field -> (now, sheet) for an existing product."""
    t = p['template']
    diff = {}
    compare = [
        ('name', t.name or '', p['name']),
        ('size', t.pos_retail_size or '', p['size']),
        ('brand', (t.brand_id.name or ''), p['brand']),
        ('price', t.list_price, p['price']),
        ('min', t.minimum_selling_price, p['minimum']),
        ('MRP', t.mrp, p['mrp']),
        ('category', t.categ_id.complete_name or '', ' / '.join(p['categ'])),
        ('till section', ' > '.join(n for n in (t.pos_categ_ids[:1].parent_id.name, t.pos_categ_ids[:1].name) if n)
         if t.pos_categ_ids else '', section_label(p['categ'])),
    ]
    if p['code']:
        compare += [('barcode', t.barcode or '', p['code']), ('internal ref', t.default_code or '', p['code'])]
    for field, now, want in compare:
        if isinstance(want, float):
            same = abs((now or 0.0) - want) <= 0.001
        elif field == 'name':
            same = clean(now) == clean(want)       # spacing alone is not a rename
        elif field in ('size', 'brand'):
            same = slug(now) == slug(want)         # nor is a change of capitals
        else:
            same = now == want
        if not same:
            diff[field] = (f"{now:g}", f"{want:g}") if isinstance(want, float) else (now or '-', want or '-')
    if any(abs(t.with_company(c).standard_price - p['cost']) > 0.001 for c in companies):
        diff['cost'] = (f"{t.with_company(branch).standard_price:g}", f"{p['cost']:g}")
    if p['unit'] and t.uom_id.name != p['unit'] and untouched(t):
        diff['unit'] = (t.uom_id.name, p['unit'])
    if t.pos_retail_branch_ids != branch:
        diff['sold at'] = (', '.join(t.pos_retail_branch_ids.mapped('name')) or 'every branch', branch.name)
    if not (t.available_in_pos and t.sale_ok and t.purchase_ok and t.is_storable and not t.company_id):
        diff['settings'] = ('off', 'sellable, stocked, shared')
    return diff


# ------------------------------------------------------------------- report
report_rows = {}
for tab in work_tabs:
    products = parsed.get(tab, [])
    new = [p for p in products if not p['template']]
    changed = [(p, differences(p)) for p in products if p['template']]
    changed = [(p, d) for p, d in changed if d]
    unchanged = len(products) - len(new) - len(changed)
    stock_load = [p for p in products if p['qty'] and (FORCE_STOCK or not p['template']
                                                       or p['template'].id not in already_counted)]
    report_rows[tab] = stock_load
    kept = sum(1 for p in products if p.get('kept'))
    print(f"\n{'-' * 78}\n{tab}: {len(products)} products  |  {len(new)} new, {len(changed)} to update, "
          f"{unchanged} unchanged  |  {sum(1 for p in products if not p['code'])} without a code"
          + (f"  |  {kept} kept after a change in the sheet" if kept else ""))
    print(f"  categories:    {sorted({' / '.join(p['categ']) for p in products})}")
    print(f"  till sections: {sorted({section_label(p['categ']) for p in products})}")
    print(f"  opening stock to load: {len(stock_load)} products, {sum(p['qty'] for p in stock_load):g}"
          + (f" (units: {sorted({p['unit'] for p in products if p['unit']})})" if any(p['unit'] for p in products) else ""))
    if new:
        print(f"  NEW ({len(new)}):")
        for p in new[:30]:
            print(f"    + row {p['row']} {p['code'] or '(no code)'} {p['name']} | {p['size'] or '-'} | "
                  f"{p['brand'] or '-'} | price {p['price']:g} | qty {p['qty'] if p['qty'] is not None else '-'}"
                  f"{' ' + p['unit'] if p['unit'] else ''}")
        if len(new) > 30:
            print(f"    ... and {len(new) - 30} more")
    if changed:
        fields = {}
        for _p, d in changed:
            for f in d:
                fields[f] = fields.get(f, 0) + 1
        print(f"  TO UPDATE ({len(changed)}): " + ', '.join(f"{f} x{n}" for f, n in sorted(fields.items())))
        for p, d in changed[:30]:
            print(f"    ~ row {p['row']} {p['code'] or '(no code)'} {p['name']}"
                  + (" [same product, row changed]" if p.get('kept') else "") + ": "
                  + '; '.join(f"{f} {a} -> {b}" for f, (a, b) in d.items()))
        if len(changed) > 30:
            print(f"    ... and {len(changed) - 30} more")
    on_hand = lambda t: t.product_variant_id.with_company(branch).with_context(
        warehouse_id=warehouse.id).qty_available
    twins = [p for p in products if p.get('retire') and not p.get('replaces')]
    if twins:
        print(f"  SECOND COPY FROM AN EARLIER RUN ({len(twins)}) -- the original keeps its stock and "
              f"history, the copy is archived and its opening stock zeroed:")
        for p in twins[:30]:
            copy = p['retire'][0]
            print(f"    x row {p['row']} {p['code'] or '(no code)'} {p['name']}: copy id {copy.id} "
                  f"(on hand {on_hand(copy):g}), original id {p['template'].id} '{p['template'].name}' "
                  f"(on hand {on_hand(p['template']):g})")
        if len(twins) > 30:
            print(f"    ... and {len(twins) - 30} more")
    replaced = [p for p in products if p.get('replaces')]
    if replaced:
        print(f"  WRONG UNIT ({len(replaced)}) -- counted but never sold; replaced by a new product in the "
              f"sheet's unit, counted from the sheet:")
        for p in replaced:
            old = p['replaces']
            print(f"    x row {p['row']} {p['name']}: id {old.id} had {on_hand(old):g} {old.uom_id.name} -> "
                  f"{p['qty'] if p['qty'] is not None else 0:g} {p['unit']}")
    out = []
    for p in products:
        if p['qty'] and p in stock_load:
            continue
        on_hand = 0.0
        if p['template']:
            on_hand = p['template'].product_variant_id.with_company(branch).with_context(
                warehouse_id=warehouse.id).qty_available
        if on_hand > 0:
            continue
        if p['qty'] is None:
            why = "quantity empty in the sheet"
        elif p['qty'] <= 0:
            why = "the sheet says 0"
        else:
            why = f"sheet says {p['qty']:g}; counted before, sold or adjusted since"
        out.append(f"row {p['row']} {p['name']}: {why}")
    if out:
        print(f"  OUT OF STOCK at the till after this run ({len(out)}):")
        for line in out:
            print(f"    - {line}")
    if tab_no_price.get(tab):
        print(f"  NO SALES PRICE ({len(tab_no_price[tab])}) -- the till won't sell these until the sheet "
              f"has one: rows {row_ranges(tab_no_price[tab])}")
    for n in tab_notes.get(tab, []):
        print(f"  note: {n}")
    groups = {}
    for reason, rownum, name in skipped.get(tab, []):
        groups.setdefault(reason, []).append(f"{rownum} {name}")
    for reason, items in groups.items():
        print(f"  skipped -- {reason} ({len(items)}): rows {', '.join(items[:6])}"
              + (f" ... +{len(items) - 6}" if len(items) > 6 else ""))

retiring = {r.id for ps in parsed.values() for p in ps for r in p.get('retire', [])}
gone = Template.browse([rid for name, rid in imported_ids.items()
                        if rid not in claimed and rid not in retiring]).exists().filtered('active')
# Gone from the sheet. Never sold or bought (nothing but the opening count): it
# is archived with that stock zeroed, or the till keeps offering a line the
# shop deleted -- often a row renamed so far it no longer resembles itself,
# whose corrected row is already a product of its own. With sales it stays,
# for the shop to decide. Only departments imported in this run, and never a
# quarter or more of one: that is a renamed tab or a bad download, not edits.
work_depts = {p['categ'][0].lower() for tab in work_tabs for p in parsed.get(tab, [])}
imported_active = Template.browse(list(imported)).exists().filtered('active')
to_archive, kept_gone = Template.browse(), Template.browse()
for dept in sorted({department(t) for t in gone}):
    in_dept = gone.filtered(lambda t: department(t) == dept)
    if dept not in work_depts:
        kept_gone |= in_dept
        continue
    unsold = in_dept.filtered(only_counted)
    total = len(imported_active.filtered(lambda t: department(t) == dept))
    if unsold and len(unsold) >= 0.25 * total:
        print(f"\nNOT archiving {len(unsold)} of {total} {dept} products missing from the sheet -- too many "
              f"at once; check the tab was not renamed or emptied.")
        kept_gone |= in_dept
        continue
    to_archive |= unsold
    kept_gone |= in_dept - unsold
if to_archive:
    print(f"\nNO LONGER IN THE SHEET, never sold ({len(to_archive)}) -- "
          + ("archived, stock zeroed" if APPLY else "will be archived, stock zeroed") + ": "
          + ', '.join(f"{t.name} {t.pos_retail_size or ''}".strip() for t in to_archive.sorted('id')[:40])
          + (f" ... +{len(to_archive) - 40}" if len(to_archive) > 40 else ""))
if kept_gone:
    print(f"\nNO LONGER IN THE SHEET, with sales or not imported now ({len(kept_gone)}) -- left as they are; "
          f"archive them by hand if unwanted: {', '.join(kept_gone.mapped('name')[:20])}")

if not APPLY:
    print("\n" + "=" * 78)
    print("CHECK ONLY -- nothing was written. Run again with APPLY=1 to update.")
    print("=" * 78)
else:
    Category = env['product.category'].sudo()
    PosCategory = env['pos.category'].sudo()
    Brand = env['product.brand'].sudo()
    caches = {'categ': {}, 'section': {}, 'brand': {}, 'uom': {}}

    def category(path):
        key = tuple(path)
        if key not in caches['categ']:
            parent = Category.browse()
            for name in path:
                parent = Category.search([('name', '=', name), ('parent_id', '=', parent.id or False)], limit=1) \
                    or Category.create({'name': name, 'parent_id': parent.id or False})
            caches['categ'][key] = parent
        return caches['categ'][key]

    def section(categ):
        top, sub = till_section(categ)
        if (top, sub) not in caches['section']:
            parent = PosCategory.search([('name', '=ilike', top), ('parent_id', '=', False)], limit=1) \
                or PosCategory.create({'name': top})
            record = parent
            if sub:
                record = PosCategory.search([('name', '=ilike', sub), ('parent_id', '=', parent.id)], limit=1) \
                    or PosCategory.create({'name': sub, 'parent_id': parent.id,
                                           'sequence': SECTION_ORDER.index(sub) if sub in SECTION_ORDER else 99})
            caches['section'][(top, sub)] = record
        return caches['section'][(top, sub)]

    def brand(name):
        if not name:
            return False
        if name.lower() not in caches['brand']:
            caches['brand'][name.lower()] = Brand.with_context(active_test=False).search(
                [('name', '=ilike', name)], limit=1) or Brand.create({'name': name})
        return caches['brand'][name.lower()].id

    def uom(unit):
        if unit not in caches['uom']:
            found = env['uom.uom'].sudo().search([('name', '=', unit)], limit=1) if unit else None
            caches['uom'][unit] = found or env.ref('uom.product_uom_unit')
        return caches['uom'][unit]

    def full_vals(p):
        vals = {
            'name': p['name'], 'pos_retail_size': p['size'] or False, 'brand_id': brand(p['brand']),
            'list_price': p['price'], 'minimum_selling_price': p['minimum'], 'mrp': p['mrp'],
            'categ_id': category(p['categ']).id, 'pos_categ_ids': [(6, 0, section(p['categ']).ids)],
            'uom_id': uom(p['unit']).id,
            'pos_retail_branch_ids': [(6, 0, branch.ids)], 'type': 'consu', 'is_storable': True,
            'available_in_pos': True, 'sale_ok': True, 'purchase_ok': True, 'company_id': False,
        }
        if p['code']:
            vals.update(default_code=p['code'], barcode=p['code'])
        return vals

    FIELD_VALS = {  # report field -> the vals that set it
        'name': ['name'], 'size': ['pos_retail_size'], 'brand': ['brand_id'], 'price': ['list_price'],
        'min': ['minimum_selling_price'], 'MRP': ['mrp'], 'category': ['categ_id'],
        'till section': ['pos_categ_ids'], 'barcode': ['barcode'], 'internal ref': ['default_code'],
        'sold at': ['pos_retail_branch_ids'], 'unit': ['uom_id'],
        'settings': ['type', 'is_storable', 'available_in_pos', 'sale_ok', 'purchase_ok', 'company_id'],
    }
    def retire(copy):
        """Zero the copy's stock and archive it, freeing its barcode for the original."""
        variant = copy.product_variant_id
        quants = env['stock.quant'].sudo().search([('product_id', '=', variant.id),
                                                   ('location_id.usage', '=', 'internal'), ('quantity', '!=', 0)])
        if quants:
            zero = env['stock.quant'].sudo().with_company(branch).with_context(inventory_mode=True).create([
                {'product_id': variant.id, 'location_id': q.location_id.id, 'inventory_quantity': 0}
                for q in quants])
            if isinstance(zero.action_apply_inventory(), dict):
                raise SystemExit(f"Odoo asked to resolve an inventory conflict on {copy.name}; nothing was "
                                 f"saved. Check Inventory > Physical Inventory for this branch.")
        copy.write({'barcode': False, 'default_code': False, 'active': False})

    def remember(p):
        """Point the row's key at its product, so the next run finds it directly."""
        if p['code'] or imported_ids.get(p['key']) == p['template'].id:
            return
        record = IMD.search([('module', '=', XMLID_MODULE), ('name', '=', p['key'])], limit=1)
        if record:
            record.write({'model': 'product.template', 'res_id': p['template'].id})
        else:
            IMD.create({'module': XMLID_MODULE, 'name': p['key'], 'model': 'product.template',
                        'res_id': p['template'].id, 'noupdate': True})
        imported_ids[p['key']] = p['template'].id

    started = time.monotonic()
    print("\nimporting... (each tab is saved when it finishes; don't interrupt)", flush=True)
    for tab in work_tabs:
        products = parsed.get(tab, [])
        with env.cr.savepoint():
            created = updated = 0
            for p in products:
                for copy in p.get('retire', []):
                    retire(copy)
            for p in products:
                t = p['template']
                if t:
                    remember(p)
                    diff = differences(p)
                    vals = full_vals(p)
                    write = {k: vals[k] for f in diff for k in FIELD_VALS.get(f, []) if k in vals}
                    if write:
                        t.write(write)
                        updated += 1
                    if t.uom_id != uom(p['unit']) and p['unit']:
                        print(f"  note: {p['name']} is sold per {t.uom_id.name} in Odoo, per {p['unit']} in the "
                              f"sheet; unit left unchanged, it already has stock or sales in {t.uom_id.name}")
            new = [p for p in products if not p['template']]
            if new:
                records = Template.create([full_vals(p) for p in new])
                taken = set(imported_ids)
                xmlids = []
                for p, record in zip(new, records):
                    name = f"cat_{slug(p['code'])}" if p['code'] else p['key']
                    p['template'] = record
                    if imported_ids.get(name) in retiring:  # the replaced product's key
                        IMD.search([('module', '=', XMLID_MODULE), ('name', '=', name)]).write(
                            {'res_id': record.id})
                        imported_ids[name] = record.id
                        continue
                    while name in taken:
                        name += '_x'
                    taken.add(name)
                    xmlids.append({'module': XMLID_MODULE, 'name': name, 'model': 'product.template',
                                   'res_id': record.id, 'noupdate': True})
                    p['template'] = record
                IMD.create(xmlids)
                created = len(new)
            # cost is company-dependent: one write per company per distinct value
            for company in companies:
                groups = {}
                for p in products:
                    if abs(p['template'].with_company(company).standard_price - p['cost']) > 0.001:
                        groups.setdefault(p['cost'], Template.browse())
                        groups[p['cost']] |= p['template']
                for value, records in groups.items():
                    records.with_company(company).write({'standard_price': value})
            # opening stock, re-checked inside the transaction
            wanted = {p['template']: p['qty'] for p in products if p['qty']}
            done = set() if FORCE_STOCK else counted(Template.browse([t.id for t in wanted]))
            load = {t: q for t, q in wanted.items() if t.id not in done}
            if load:
                quants = env['stock.quant'].sudo().with_company(branch).with_context(inventory_mode=True).create([{
                    'product_id': t.product_variant_id.id, 'location_id': warehouse.lot_stock_id.id,
                    'inventory_quantity': q} for t, q in load.items()])
                if isinstance(quants.action_apply_inventory(), dict):
                    raise SystemExit("Odoo asked to resolve an inventory conflict; nothing was saved. "
                                     "Check Inventory > Physical Inventory for this branch.")
        env.cr.commit()
        retired = sum(len(p.get('retire', [])) for p in products)
        print(f"  {tab}: created {created}, updated {updated}, opening stock on {len(load)} products "
              f"({sum(load.values()):g})" + (f", copies/wrong-unit products archived {retired}" if retired else "")
              + f"  [{time.monotonic() - started:.0f}s]", flush=True)
    if to_archive:
        with env.cr.savepoint():
            for t in to_archive:
                retire(t)
        env.cr.commit()
        print(f"  no longer in the sheet: archived {len(to_archive)} never-sold products, stock zeroed",
              flush=True)
    # An earlier run filed rows under a marker row ("NEW") as if it were a
    # category; once its products have moved back, the empty leftovers go.
    for marker in MARKER_HEADINGS:
        empty = [c for c in Category.search([('name', '=ilike', marker), ('parent_id', '!=', False)])
                 if not c.child_id and not Template.search_count([('categ_id', '=', c.id)])]
        empty += [s for s in PosCategory.search([('name', '=ilike', marker), ('parent_id', '!=', False)])
                  if not s.child_ids and not Template.search_count([('pos_categ_ids', 'in', s.ids)])]
        for record in empty:
            label = f"{record.parent_id.name} > {record.name}"
            try:
                with env.cr.savepoint():
                    record.unlink()
                print(f"  removed the empty {record._description.lower()} {label}")
            except UserError as e:  # e.g. a till section while a register session is open
                print(f"  left the empty {record._description.lower()} {label}: {e} -- run again later")
    # Brands that were really sizes typed into the Brand column ("32 x 25 mm"),
    # left without products once the sheet was corrected. Only names with a
    # digit: a real brand ("Master") may sit in the size column of a bad row.
    sizes = {slug(p['size']) for ps in parsed.values() for p in ps if p['size']}
    brands_used = {slug(p['brand']) for ps in parsed.values() for p in ps if p['brand']}
    for b in Brand.with_context(active_test=False).search([]):
        if (re.search(r'\d', b.name) and slug(b.name) in sizes and slug(b.name) not in brands_used
                and not Template.search_count([('brand_id', '=', b.id)])):
            try:
                with env.cr.savepoint():
                    name = b.name
                    b.unlink()
                print(f"  removed the brand {name!r} (a size typed in the Brand column, no products left)")
            except UserError as e:
                print(f"  left the brand {b.name!r}: {e}")
    env.cr.commit()
    print("=" * 78)
    print("Import finished. Reload the till pages.")
    print("=" * 78)
